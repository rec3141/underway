"""The report as a .docx, written into Amundsen Science's cruise-report template.

The template (``Cruise Report Template_2026.docx`` on the share) is filled in
place so its logo, styles, survey link and CFI section survive: the
"save under a new name" notice and the grey guidance lines (style
``Subtitle``) are removed, the title block and each section's
``[Text goes here]`` are replaced with the participant's text, and the
generated narrative, tables and figures are inserted after that text in the
section the participant chose.

A table wider than ``WIDE_COLUMNS`` gets landscape pages of its own
(section breaks either side); tables span the text width.

Captions use Word SEQ fields, so numbering stays right when the participant
adds their own tables and figures afterwards (Word updates them on print or
F9).
"""

from __future__ import annotations

import copy
import io
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from . import tables as T
from .config import TEMPLATE

SECTIONS = {
    "intro": "Introduction & Objectives",
    "methods": "Methodology",
    "results": "Preliminary Results",
    "references": "References",
    "recommendations": "Recommendations",
}
# A table with more columns than this goes on landscape pages of its own.
WIDE_COLUMNS = 8

CFI = {"publications": "Publications", "presentations": "Presentations",
       "in_progress": "In-Progress"}


def _para_after(anchor, doc, text="", style=None):
    p = doc.add_paragraph(text, style=style)
    anchor.addnext(p._p)
    return p


def _move_after(anchor, el):
    anchor.addnext(el)
    return el


def _caption(doc, anchor, kind: str, text: str):
    """'Table 1. text' with the number as a SEQ field."""
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(f"{kind} ")
    r.bold = True
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" SEQ {kind} \\* ARABIC ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    b = OxmlElement("w:b")
    rpr.append(b)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = str(_caption.counter.setdefault(kind, 0) + 1)
    _caption.counter[kind] += 1
    run.append(t)
    fld.append(run)
    p._p.append(fld)
    r2 = p.add_run(". ")
    r2.bold = True
    p.add_run(text)
    anchor.addnext(p._p)
    return p._p


_caption.counter = {}


def _borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "A6A6A6")
        borders.append(el)
    tbl_pr.append(borders)


def _section_break(doc, anchor, landscape: bool):
    """A paragraph that ends the current section; the section it ends is
    landscape or portrait as asked, with the template's margins."""
    base = doc.sections[0]._sectPr
    sect = copy.deepcopy(base)
    for tag in ("w:headerReference", "w:footerReference", "w:titlePg"):
        for el in sect.findall(qn(tag)):
            sect.remove(el)
    pg = sect.find(qn("w:pgSz"))
    w, hh = int(pg.get(qn("w:w"))), int(pg.get(qn("w:h")))
    short, long_ = sorted((w, hh))
    pg.set(qn("w:w"), str(long_ if landscape else short))
    pg.set(qn("w:h"), str(short if landscape else long_))
    if landscape:
        pg.set(qn("w:orient"), "landscape")
    elif pg.get(qn("w:orient")):
        del pg.attrib[qn("w:orient")]
    typ = sect.find(qn("w:type"))
    if typ is None:
        typ = OxmlElement("w:type")
        sect.insert(0, typ)
    typ.set(qn("w:val"), "nextPage")
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().append(sect)
    anchor.addnext(p._p)
    return p._p


def _full_width(table):
    tbl_pr = table._tbl.tblPr
    w = tbl_pr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW")
        tbl_pr.append(w)
    w.set(qn("w:type"), "pct")
    w.set(qn("w:w"), "5000")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tbl_pr.append(layout)


def _table(doc, anchor, tab: dict, *, font_pt: float = 8):
    cols, body = tab["columns"], tab["body"]
    t = doc.add_table(rows=1 + len(body), cols=len(cols))
    _borders(t)
    _full_width(t)
    for j, c in enumerate(cols):
        cell = t.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(c["label"] + (f" ({c['unit']})" if c.get("unit") else ""))
        run.bold = True
        run.font.size = Pt(font_pt)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "EDEDED")
        cell._tc.get_or_add_tcPr().append(shd)
    # Repeat the header row across pages.
    trpr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    trpr.append(hdr)
    for i, row in enumerate(body, start=1):
        cells = t.rows[i].cells
        for j, (v, c) in enumerate(zip(row, cols)):
            cells[j].text = ""
            para = cells[j].paragraphs[0]
            run = para.add_run(T.fmt(v, c))
            run.font.size = Pt(font_pt)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
    return _move_after(anchor, t._tbl)


def _picture(doc, anchor, png: bytes, width_cm: float = 16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(io.BytesIO(png), width=Cm(width_cm))
    return _move_after(anchor, p._p)


def _set_text(par, text: str):
    """Replace a paragraph's runs with plain text, keeping its style."""
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    for h in par._p.findall(qn("w:hyperlink")):
        par._p.remove(h)
    par.add_run(text)


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]


def build(report: dict, content: dict) -> bytes:
    """``report``: the participant's text and choices; ``content``: generated parts.

    content = {
      "narratives": [str, ...],                         # station conditions, in order
      "blocks": [{"section", "kind": "table"|"figure", "caption", "table"|"png"}],
    }
    """
    _caption.counter = {}
    doc = docx.Document(str(TEMPLATE))
    body = doc.element.body

    # The "save under a new name" notice box.
    for tbl in list(body.findall(qn("w:tbl"))):
        if "Important Action Required" in "".join(tbl.itertext()):
            body.remove(tbl)

    h = report.get("header", {})
    leaders = [x for x in h.get("leaders", []) if x.get("name")]
    people = [x for x in h.get("participants", []) if x.get("name")]
    affs: list[str] = []

    def sup(aff):
        if not aff:
            return ""
        if aff not in affs:
            affs.append(aff)
        return str(affs.index(aff) + 1)

    heading_of: dict[str, object] = {}
    for p in doc.paragraphs:
        st, txt = p.style.name, p.text.strip()
        if st == "Heading 2":
            _set_text(p, h.get("title") or "Brief Project Title")
        elif st == "No Spacing" and txt.startswith("Project leaders"):
            _person_line(p, "Project leaders", leaders, sup, emails=True)
        elif st == "No Spacing" and txt.startswith("Cruise participants"):
            leg = report.get("leg_label") or "[number]"
            _person_line(p, f"Cruise participants {leg}", people, sup)
        elif st == "Heading 3":
            heading_of[txt] = p

    # Affiliation lines: replace the template's two with ours.
    aff_pars = [p for p in doc.paragraphs if p.style.name == "No Spacing"
                and p.text.strip().lstrip("0123456789").startswith("Affiliation")]
    if aff_pars:
        anchor = aff_pars[-1]._p
        for p in aff_pars[:-1]:
            p._p.getparent().remove(p._p)
        last = aff_pars[-1]
        _set_text(last, "")
        lines = affs or ["[Institute/University name]"]
        for i, a in enumerate(lines):
            p = last if i == 0 else _para_after(anchor, doc, "", "No Spacing")
            r = p.add_run(str(i + 1))
            r.font.superscript = True
            r.italic = True
            p.add_run(a).italic = True
            anchor = p._p

    # Guidance lines go; the Survey section's note stays (it is the template's text).
    for p in list(doc.paragraphs):
        if p.style.name == "Subtitle" and "User feedback is vital" not in p.text:
            p._p.getparent().remove(p._p)

    text = report.get("text", {})
    for key, heading in SECTIONS.items():
        hp = heading_of.get(heading)
        if hp is None:
            continue
        # The section's placeholder paragraphs, up to the next heading.
        placeholders = []
        el = hp._p.getnext()
        while el is not None and el.tag == qn("w:p"):
            par = docx.text.paragraph.Paragraph(el, hp._parent)
            if par.style.name.startswith("Heading"):
                break
            if re.match(r"^\[(Text goes here|Reference \d+)\]$", par.text.strip()):
                placeholders.append(par)
            el = el.getnext()
        anchor = hp._p
        for par in placeholders:
            par._p.getparent().remove(par._p)
        if key == "references":
            for line in [x.strip() for x in (text.get("references") or "").splitlines() if x.strip()]:
                anchor = _para_after(anchor, doc, line, "List Paragraph")._p
            if not (text.get("references") or "").strip():
                anchor = _para_after(anchor, doc, "[Reference 1]", "List Paragraph")._p
            continue
        chunks = _paragraphs(text.get(key, ""))
        for chunk in chunks or ["[Text goes here]"]:
            anchor = _para_after(anchor, doc, chunk, "Normal")._p
        if key == "methods" and content.get("narratives"):
            p = _para_after(anchor, doc, "", "Normal")
            p.add_run("Conditions at the stations").bold = True
            anchor = p._p
            for n in content["narratives"]:
                anchor = _para_after(anchor, doc, n, "Normal")._p
        for b in content.get("blocks", []):
            if b.get("section", "methods") != key:
                continue
            if b["kind"] == "table":
                wide = len(b["table"]["columns"]) > WIDE_COLUMNS
                if wide:
                    anchor = _section_break(doc, anchor, landscape=False)
                anchor = _caption(doc, anchor, "Table", b.get("caption") or b["table"].get("title", ""))
                anchor = _table(doc, anchor, b["table"])
                if wide:
                    anchor = _section_break(doc, anchor, landscape=True)
                else:
                    anchor = _para_after(anchor, doc, "", "Normal")._p
            else:
                anchor = _picture(doc, anchor, b["png"])
                anchor = _caption(doc, anchor, "Figure", b.get("caption", ""))

    # CFI outputs: fill each bullet after its bold lead-in.
    for p in doc.paragraphs:
        if p.style.name != "List Paragraph":
            continue
        for key, lead in CFI.items():
            if p.text.strip().startswith(lead) and text.get(key):
                for r in p.runs[1:]:
                    r._r.getparent().remove(r._r)
                p.add_run(": " + text[key].strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _person_line(p, lead: str, people: list[dict], sup, emails=False):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for h in p._p.findall(qn("w:hyperlink")):
        p._p.remove(h)
    p.add_run(f"{lead}: ").bold = True
    if not people:
        p.add_run("Name")
        return
    for i, x in enumerate(people):
        if i:
            p.add_run(", ")
        p.add_run(x["name"])
        n = sup(x.get("affiliation"))
        if n:
            p.add_run(n).font.superscript = True
        if emails and x.get("email"):
            p.add_run(f" ({x['email']})")


def word_count(report: dict, narratives: list[str]) -> int:
    text = " ".join(str(v) for v in (report.get("text") or {}).values()) + " " + " ".join(narratives)
    return len(re.findall(r"\b\w[\w'’-]*\b", text))

