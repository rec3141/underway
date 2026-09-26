"""Pure-function checks that need no shares mounted."""

import pandas as pd

from cruisereport import activities, conditions, logsheets, rosette


def test_activity_spellings_fold_into_groups():
    cases = {
        "TM-Rosette P+A": "tm_rosette", "TM-CTD": "tm_rosette", "CTD (5 bottles only)": "rosette",
        "Classic-Rosette Limited": "rosette", "Box Core - BIO": "box_core",
        "PC or GGC": "gravity_piston", "IKMT ou Beam Trawl": "ikmt_beam",
        "Monster or Hydrobios": "plankton_nets", "Transit + MVP to Basic+ CS2-1": "mvp",
        "Transit to Qaanaq": "transit", "Deploy helicpopter": "helicopter",
        "CTD-Rosette (erreur, à enlever)": "void", "Something new": "other",
    }
    for text, group in cases.items():
        assert activities.group_of(text) == group, text


def test_rosette_ice_terms():
    assert rosette.ice_english("1/10 Eau libre") == "open water (under 1/10)"
    assert rosette.ice_english("4-6/10 Banquise lâche") == "4-6/10 open drift ice"
    assert rosette.ice_english("7") is None          # a bare number: no scale stated
    assert rosette.ice_english("-99") is None


def test_sun_elevation_polar_night_and_noon():
    # Midsummer local noon at 75°N: about 90 - 75 + 23.4 degrees.
    assert abs(conditions.sun_elevation("2026-06-21T17:00:00", 75.0, -75.0) - 38.4) < 1.0
    assert conditions.sun_elevation("2026-12-21T05:00:00", 75.0, -75.0) < -30


def test_range_formatting():
    assert conditions._range([-0.4, -0.2], 0, "°") == "0°"
    assert conditions._range([1.0, 4.2], 0, "°") == "1 to 4°"
    assert conditions._range([-1.25, 0.5], 1, "°C") == "−1.2 to 0.5 °C"


def test_logsheet_roles_from_names_and_values():
    df = pd.DataFrame({"Unique Identifier": ["AMD2603-010", "AMD2603-011"],
                       "Station Name": ["CardS-3", "CardS-3"],
                       "Date & Time (UTC)": ["2026/09/06 13:21:00", "2026/09/06 14:59:00"],
                       "Latitude": [76.3, 76.3], "Longitude": [-89.2, -89.2]})
    roles = logsheets.guess_roles(df)
    assert roles["label"] == "Unique Identifier"
    assert roles["station"] == "Station Name"
    assert roles["datetime"] == "Date & Time (UTC)"
    assert roles["lat"] == "Latitude" and roles["lon"] == "Longitude"


def test_digitize_levels_and_ditto(tmp_path, monkeypatch):
    from cruisereport import digitize

    assert [digitize.level(x) for x in (3, -1, 100, 90, 75, 50, 10)] == [3, -1, 3, 2, 1, 0, -1]
    assert digitize.colour(3) != digitize.colour(-1)
    for mark in ('"', "〃", "↓", "=", "do", "Ditto"):
        assert digitize.DITTO.match(mark), mark
    assert not digitize.DITTO.match("079")

    monkeypatch.setattr(digitize, "STATE_DIR", tmp_path)
    cell = lambda t, c=3: {"t": t, "c": c}  # noqa: E731
    doc = digitize.save("page.jpg", b"jpeg", {"tables": [{"title": "log", "columns": ["ST", "CAST", "BOT"], "rows": [
        [cell("LAS-2"), cell("079"), cell("1")],
        [cell(""), cell("↓"), cell("2")],
        [cell("CONTROL"), cell(""), cell("HOT WATER")],
    ]}], "notes": []})
    cols, rows = digitize.rows_for_logsheet(doc["id"], 0, True)
    assert rows[1] == {"ST": "LAS-2", "CAST": "079", "BOT": "2"}          # continues the record
    assert rows[2]["CAST"] == ""                                          # a row of its own
    digitize.edit(doc["id"], 0, 2, 1, "080")
    assert digitize.load(doc["id"])["tables"][0]["rows"][2][1] == {"t": "080", "c": 3, "edited": True}


def test_digitize_queue_runs_fails_and_resumes(tmp_path, monkeypatch):
    from cruisereport import digitize

    monkeypatch.setattr(digitize, "STATE_DIR", tmp_path)
    monkeypatch.setattr(digitize, "_workers", [])
    monkeypatch.setattr(digitize, "_jobs", digitize.queue.Queue())
    answers = {"good.jpg": {"tables": [{"title": "t", "columns": ["A"], "rows": [[{"t": "1", "c": "likely"}]]}],
                            "notes": [], "model": "m", "usage": {}, "seconds": 1}}

    def fake(jpeg):
        name = jpeg.decode()
        if name not in answers:
            raise RuntimeError("model unavailable")
        return answers[name]
    monkeypatch.setattr(digitize, "transcribe_page", fake)

    # A page left "working" by a restart is taken up again on start.
    stale = digitize.submit("good.jpg", b"good.jpg")
    doc = digitize.load(stale["id"])
    doc["status"] = "working"
    digitize._write(stale["id"], doc)
    digitize._jobs = digitize.queue.Queue()          # the restart's empty in-memory queue

    digitize.start(workers=1)
    good = digitize.submit("good.jpg", b"good.jpg")
    bad = digitize.submit("bad.jpg", b"bad.jpg")
    digitize._jobs.join()
    assert digitize.load(stale["id"])["status"] == "done"
    done = digitize.load(good["id"])
    assert done["status"] == "done" and done["tables"][0]["rows"][0][0] == {"t": "1", "c": 1}
    failed = digitize.load(bad["id"])
    assert failed["status"] == "failed" and "model unavailable" in failed["error"]


def test_ticked_logs_name_the_teams_bottles():
    from cruisereport import tables

    logs = {"log:a:s": {"name": "eDNA", "roles": {"bottle": "BOT"}, "rows": [
        {"BOT": "1", "_op": "AMD2603-230"}, {"BOT": " 2 ", "_op": "AMD2603-230"},
        {"BOT": "HOT WATER", "_op": "AMD2603-230"}, {"BOT": "3", "_op": None}]}}
    used = [{"id": "a", "sheet": "s", "use": True}]
    assert tables.logged_bottles(logs, used) == {("AMD2603-230", 1): "eDNA", ("AMD2603-230", 2): "eDNA"}
    assert tables.logged_bottles(logs, [{**used[0], "use": False}]) == {}
    b = {"label": "AMD2603-230", "bottle": 2, "draws": {}}
    assert tables._team_bottle(b, [], tables.logged_bottles(logs, used))
    assert not tables._team_bottle({**b, "bottle": 5}, [], tables.logged_bottles(logs, used))


def test_digitized_tables_grow(tmp_path, monkeypatch):
    import pytest
    from cruisereport import digitize

    monkeypatch.setattr(digitize, "STATE_DIR", tmp_path)
    doc = digitize.save("p.jpg", b"j", {"tables": [{"title": "t", "columns": ["A", "B"],
                                                    "rows": [[{"t": "1", "c": 2}, {"t": "2", "c": 2}]]}], "notes": []})
    digitize.grow(doc["id"], 0, "row")
    t = digitize.grow(doc["id"], 0, "col")["tables"][0]
    assert t["columns"] == ["A", "B", "column 3"]
    assert [len(r) for r in t["rows"]] == [3, 3]
    assert t["rows"][1][0] == {"t": "", "c": 3, "edited": True}
    with pytest.raises(ValueError):
        digitize.grow(doc["id"], 0, "sideways")


def test_log_columns_by_header_and_role():
    from cruisereport import tables

    a = {"log": {"ST": "LAS-2", "Depth ": "529", "_op": "x"}, "lg": {"name": "p1", "roles": {"station": "ST"}}}
    b = {"log": {"STATION": "ES2", "depth": "30"}, "lg": {"name": "p3", "roles": {"station": "STATION"}}}
    assert [tables._value(r, "log.DEPTH") for r in (a, b)] == ["529", "30"]      # case and spacing
    assert [tables._value(r, "logrole.station") for r in (a, b)] == ["LAS-2", "ES2"]
    assert tables._value(a, "logmeta.log") == "p1"
    assert tables._value(a, "log.CAST") is None


def test_uploaded_logsheet_edits_matches_and_exports(tmp_path, monkeypatch):
    import io

    import openpyxl
    import pytest

    monkeypatch.setattr(logsheets, "STATE_DIR", tmp_path)
    csv = b"Station,Depth,Sample ID\nBay Fiord,10,0012\nBay Fiord,20,0013\n"
    meta = logsheets.save_upload("net.csv", csv)
    ident, sheet = meta["id"], "csv"

    logsheets.edit(ident, sheet, 0, 1, "12.5")
    logsheets.edit(ident, sheet, 1, 2, "0014")
    logsheets.edit(ident, sheet, -1, 1, "Depth (m)")
    logsheets.grow(ident, sheet, "row")
    logsheets.grow(ident, sheet, "col")
    sh = logsheets.load(ident)["sheets"][sheet]
    assert sh["columns"] == ["Station", "Depth (m)", "Sample ID", "column 4"]
    assert sh["rows"][0]["Depth (m)"] == 12.5 and sh["rows"][1]["Sample ID"] == "0014"
    assert len(sh["rows"]) == 3 and sh["rows"][2] == dict.fromkeys(sh["columns"])
    assert sh["edited"] == [[0, 1], [1, 2]]
    with pytest.raises(ValueError):
        logsheets.edit(ident, sheet, -1, 0, "Sample ID")          # a column name already taken

    logsheets.set_match(ident, sheet, 2, "OPKEY")
    assert logsheets.load(ident)["sheets"][sheet]["manual"] == {"2": "OPKEY"}
    logsheets.set_match(ident, sheet, 2, None)
    assert logsheets.load(ident)["sheets"][sheet]["manual"] == {}

    assert logsheets.tsv(ident, sheet).splitlines()[:2] == [
        "Station\tDepth (m)\tSample ID\tcolumn 4", "Bay Fiord\t12.5\t0012\t"]
    ws = openpyxl.load_workbook(io.BytesIO(logsheets.xlsx(ident)))["csv"]
    assert ws["B2"].value == 12.5 and ws["B2"].fill.fgColor.rgb.endswith(ws["C3"].fill.fgColor.rgb[-6:])
    assert ws["A2"].fill.fill_type is None

    # A log made from a transcribed table is corrected there, not here.
    made = logsheets.save_frames("page · table 1", {"transcribed": pd.DataFrame({"A": [1]})},
                                 {"digitized": "0" * 12, "table": 0})
    with pytest.raises(ValueError):
        logsheets.edit(made["id"], "transcribed", 0, 0, "2")


def test_any_automatic_match_can_be_removed_or_replaced_by_hand(monkeypatch):
    class Op:
        def __init__(self, key, station, t):
            self.key, self.group = key, "ctd"
            self._s = {"key": key, "label": key, "station": station, "group": "ctd", "start_utc": t,
                       "lat": None, "lon": None}

        def summary(self):
            return self._s

    monkeypatch.setattr(logsheets.eventlog, "operations",
                        lambda leg: [Op("AMD2603-001", "S1", "2026-09-01T10:00:00"),
                                     Op("AMD2603-002", "S2", "2026-09-02T10:00:00")])
    monkeypatch.setattr(logsheets.ctd, "label_for_cast", lambda leg: {})
    roles = {"station": "Stn", "label": "Event"}
    rows = [{"Event": "AMD2603-001", "Stn": "S1"},                        # by label
            {"Event": "", "Stn": "S2"},                                   # by the one visit to S2
            {"Event": "AMD2603-001", "Stn": "S1", logsheets.HAND_COLUMN: logsheets.NO_MATCH},
            {"Event": "", "Stn": "S2", logsheets.HAND_COLUMN: logsheets.NO_MATCH},
            {"Event": "AMD2603-001", "Stn": "S1", logsheets.HAND_COLUMN: "AMD2603-002"}]
    got = [(r["_op"], r["_how"]) for r in logsheets.match(rows, roles, "leg")]
    assert got == [("AMD2603-001", "label"), ("AMD2603-002", "station"),
                   (None, "removed by hand"), (None, "removed by hand"), ("AMD2603-002", "by hand")]


def test_digitized_column_roles_are_kept_and_follow_a_rename(tmp_path, monkeypatch):
    from cruisereport import digitize

    monkeypatch.setattr(digitize, "STATE_DIR", tmp_path)
    doc = digitize.save("p.jpg", b"j", {"tables": [{"title": "t", "columns": ["STN", "T"],
                                                    "rows": [[{"t": "S1", "c": 2}, {"t": "10:00", "c": 2}]]}], "notes": []})
    digitize.set_roles(doc["id"], 0, {"station": "STN", "time": "T", "cast": ""})
    assert digitize.load(doc["id"])["tables"][0]["roles"] == {"station": "STN", "time": "T"}
    digitize.edit(doc["id"], 0, -1, 1, "Time (UTC)")
    assert digitize.load(doc["id"])["tables"][0]["roles"] == {"station": "STN", "time": "Time (UTC)"}
    digitize.set_roles(doc["id"], 0, None)
    assert "roles" not in digitize.load(doc["id"])["tables"][0]
